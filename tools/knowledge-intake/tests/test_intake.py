from __future__ import annotations

import hashlib
import importlib.util
import json
import os
import subprocess
import sys
import tempfile
import threading
import unittest
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from unittest.mock import patch

from knowledge_intake import sync_source, sync_sources
from knowledge_intake.common import ImportLimit, command, digest
from knowledge_intake.generation import GenerationPending, accept_native_result, generate, request_record
from knowledge_intake.sources import normalize, source_identity

HAS_DOCUMENTS = all(importlib.util.find_spec(module) for module in ("docx", "pptx", "openpyxl", "pypdfium2", "PIL", "reportlab"))


class IntakeTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name).resolve()
        self.source = self.root / "input"
        self.source.mkdir()
        self.state = self.root / "state"
        self.output = self.root / "output"

    def tearDown(self):
        self.temp.cleanup()

    def sync(self, spec=None, **kwargs):
        return sync_source(spec or str(self.source), state_root=self.state, output_root=self.output, **kwargs)

    def note(self):
        return next(self.output.glob("*/*.md"))

    def assert_spans(self, note):
        body = note.read_text(encoding="utf-8")
        metadata = json.loads(note.with_suffix(".meta.json").read_text(encoding="utf-8"))
        for span in metadata["evidence"]["spans"]:
            actual = "\n".join(body.splitlines()[span["line_start"] - 1:span["line_end"]])
            self.assertEqual(digest(actual.encode()), span["sha256"], span)
        return metadata

    def test_incremental_and_preserve_maintainer(self):
        raw = b"# HCCL\r\n\r\n910B graph replay\r\n"
        path = self.source / "hccl.md"
        path.write_bytes(raw)
        first = self.sync()
        self.assertEqual((first["status"], first["converted"], first["updated"]), ("ok", 1, 1))
        metadata = self.assert_spans(self.note())
        self.assertEqual(metadata["source"]["sha256"], hashlib.sha256(raw).hexdigest())
        self.assertEqual(len(list(self.output.rglob("*.md"))), 1, "raw Markdown assets must not be indexed a second time")
        self.assertEqual(self.sync()["converted"], 0)
        path.write_text("# HCCL\n\n910C graph replay", encoding="utf-8")
        second = self.sync()
        self.assertEqual(second["updated"], 1)
        self.assertNotEqual(first["revision"], second["revision"])
        self.note().write_text("# Curated\n\nMaintainer interpretation", encoding="utf-8")
        path.write_text("# New source", encoding="utf-8")
        third = self.sync(force=True)
        self.assertEqual(third["preserved"], ["hccl.md"])
        self.assertEqual(third["revision"], second["revision"])
        self.assertIn("Maintainer", self.note().read_text(encoding="utf-8"))

    def test_metadata_and_original_asset_edits_preserved(self):
        (self.source / "x.txt").write_text("NPU", encoding="utf-8")
        self.sync()
        self.note().with_suffix(".meta.json").write_text("{}", encoding="utf-8")
        self.assertEqual(len(self.sync()["preserved"]), 1)

    def test_limit_failure_retains_cursor_and_missing_is_not_deleted(self):
        path = self.source / "first.md"
        path.write_text("Ascend", encoding="utf-8")
        first = self.sync()
        (self.source / "large.txt").write_text("x" * 300, encoding="utf-8")
        partial = self.sync({"path": str(self.source), "limits": {"file_bytes": 100}})
        self.assertEqual(partial["status"], "partial")
        self.assertEqual(partial["revision"], first["revision"])
        self.assertEqual(partial["missing"], [])
        (self.source / "large.txt").unlink()
        path.unlink()
        result = self.sync()
        self.assertEqual(result["missing"], ["first.md"])
        self.assertTrue(self.note().exists())

    def test_scan_bound_counts_unsupported_entries(self):
        for index in range(8):
            (self.source / f"{index}.ignored").touch()
        result = self.sync({"path": str(self.source), "limits": {"scan_entries": 3}})
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["scanned_entries"], 4)
        self.assertEqual(result["converted"], 0)

    def test_html_preserves_headings_table_and_code_not_scripts(self):
        (self.source / "performance.html").write_text("<h1>Ascend throughput</h1><script>BAD_RUN()</script><table><tr><th>Device</th><th>Tokens/s</th></tr><tr><td>910B</td><td>123</td></tr></table><pre>x = 1\ny = 2</pre>", encoding="utf-8")
        self.assertEqual(self.sync()["status"], "ok")
        body = self.note().read_text(encoding="utf-8")
        self.assertIn("# Ascend throughput", body)
        self.assertIn("910B | 123", body)
        self.assertIn("x = 1\ny = 2", body)
        self.assertNotIn("BAD_RUN", body)
        self.assert_spans(self.note())

    def test_symlink_outside_source_is_not_followed(self):
        outside = self.root / "secret.md"
        outside.write_text("outside", encoding="utf-8")
        try:
            (self.source / "link.md").symlink_to(outside)
        except OSError:
            self.skipTest("host does not permit symlink creation")
        self.assertEqual(self.sync()["converted"], 0)

    def test_output_symlink_escape_rejected(self):
        (self.source / "x.md").write_text("Ascend", encoding="utf-8")
        ident = source_identity(normalize(str(self.source)))
        self.output.mkdir()
        try:
            (self.output / ident).symlink_to(self.source, target_is_directory=True)
        except OSError:
            self.skipTest("host does not permit symlink creation")
        result = self.sync()
        self.assertEqual(result["status"], "partial")
        self.assertIn("escapes", str(result["errors"]))

    def test_interrupted_pair_write_recovers_without_reconversion(self):
        (self.source / "x.md").write_text("Ascend", encoding="utf-8")
        module = __import__("knowledge_intake.sync", fromlist=["sync_source"])
        original = module.write_atomic
        def fail_metadata(path, data):
            if str(path).startswith(str(self.output)) and path.suffix == ".json":
                raise OSError("simulated interruption after body write")
            return original(path, data)
        with patch.object(module, "write_atomic", side_effect=fail_metadata):
            first = self.sync()
        self.assertEqual(first["status"], "partial")
        second = self.sync()
        self.assertEqual(second["status"], "ok")
        self.assertEqual(second["converted"], 0)
        self.assert_spans(self.note())

    def test_native_caption_revision_and_real_file_handoff_contract(self):
        image = self.source / "chart.png"
        image.write_bytes(b"test fixture image bytes; no claimed image interpretation")
        request = request_record("image_caption", "Describe", [image], ["source-ref"])
        with self.assertRaises(ValueError):
            accept_native_result(request, {"input_sha256": "wrong", "text": "fixture"})
        with self.assertRaises(ValueError):
            accept_native_result(request, {"input_sha256": request["input_sha256"], "text": "fixture", "refs": ["unknown"]})
        config = {"directory": str(self.root / "native")}
        with self.assertRaises(GenerationPending):
            generate(config, task="image_caption", text="Describe", images=[image], allowed_refs=["source-ref"])
        stored = json.loads(next((self.root / "native").glob("*.request.json")).read_text(encoding="utf-8"))
        (self.root / "native" / (stored["input_sha256"] + ".result.json")).write_text(json.dumps({"input_sha256": stored["input_sha256"], "text": "Fixture-only result", "refs": ["source-ref"]}), encoding="utf-8")
        self.assertEqual(generate(config, task="image_caption", text="Describe", images=[image], allowed_refs=["source-ref"])["text"], "Fixture-only result")
        image.write_bytes(b"changed")
        with self.assertRaises(ValueError):
            accept_native_result(request, {"input_sha256": request["input_sha256"], "text": "fixture"})

    def test_git_pinned_snapshot_ignores_uncommitted_edits(self):
        def git(*args):
            return subprocess.check_output(["git", *args], cwd=self.source, stderr=subprocess.DEVNULL).decode().strip()
        git("init", "-q")
        git("config", "user.name", "Intake Fixture")
        git("config", "user.email", "fixture@example.invalid")
        path = self.source / "README.md"
        path.write_text("# Pinned Ascend source", encoding="utf-8")
        git("add", ".")
        git("commit", "-qm", "fixture")
        revision = git("rev-parse", "HEAD")
        path.write_text("uncommitted source", encoding="utf-8")
        result = self.sync({"type": "git", "path": str(self.source)})
        self.assertEqual(result["revision"], revision)
        self.assertIn("Pinned Ascend", self.note().read_text(encoding="utf-8"))
        self.assertNotIn("uncommitted", self.note().read_text(encoding="utf-8"))
        self.assertEqual(self.sync({"type": "git", "path": str(self.source)})["converted"], 0)

    def test_url_real_transport_and_failure_retains_output(self):
        class Handler(BaseHTTPRequestHandler):
            code = 200
            def do_GET(self):
                self.send_response(self.code)
                self.send_header("Content-Type", "text/html")
                self.end_headers()
                self.wfile.write(b"<h1>NPU graph execution</h1><p>source snapshot</p>")
            def log_message(self, *args):
                pass
        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            url = f"http://127.0.0.1:{server.server_port}/document"
            first = self.sync(url)
            self.assertEqual(first["status"], "ok")
            self.assertEqual(self.sync(url)["converted"], 0)
            Handler.code = 503
            failure = self.sync(url)
            self.assertEqual(failure["revision"], first["revision"])
            self.assertEqual(failure["status"], "partial")
            self.assertIn("NPU graph", self.note().read_text(encoding="utf-8"))
        finally:
            server.shutdown()
            server.server_close()
            thread.join()

    def test_slow_url_has_absolute_source_deadline(self):
        import time
        class Handler(BaseHTTPRequestHandler):
            def do_GET(self):
                self.send_response(200)
                self.send_header("Content-Type", "text/plain")
                self.end_headers()
                try:
                    for _ in range(20):
                        self.wfile.write(b"a")
                        self.wfile.flush()
                        time.sleep(.1)
                except (BrokenPipeError, ConnectionResetError, ConnectionAbortedError):
                    pass
            def log_message(self, *args):
                pass
        server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        try:
            before = time.monotonic()
            result = self.sync({"url": f"http://127.0.0.1:{server.server_port}/slow", "limits": {"seconds": .5}})
            self.assertEqual(result["status"], "partial")
            self.assertLess(time.monotonic() - before, 1.5)
            self.assertIsNone(result["revision"])
        finally:
            server.shutdown()
            server.server_close()
            thread.join()

    def test_pr_allowlist_and_api_fixture_scope(self):
        rejected = self.sync({"type": "github_pr", "repo": "Tencent/WeKnora", "number": 1})
        self.assertEqual(rejected["status"], "partial")
        pr = {"title": "Fixture: Ascend graph", "base": {"sha": "a" * 40}, "head": {"sha": "b" * 40}, "updated_at": "2026-09-13T00:00:00Z", "state": "closed", "merged": True, "changed_files": 1, "body": "Fixture data, no real PR claim"}
        files = [{"filename": "vllm_ascend/worker.py", "status": "modified", "additions": 1, "deletions": 0, "patch": "@@ -1 +1 @@\n+graph"}]
        def api(endpoint, budget):
            if "/files?" in endpoint:
                return files
            if "/comments?" in endpoint:
                return []
            return pr
        with patch("knowledge_intake.sources.github_api", side_effect=api):
            result = self.sync({"type": "github_pr", "repo": "vllm-project/vllm-ascend", "number": 1})
        self.assertEqual(result["status"], "ok")
        text = self.note().read_text(encoding="utf-8")
        self.assertIn("API patch excerpt", text)
        self.assertIn("not execution evidence", text)
        self.assertEqual(self.assert_spans(self.note())["source"]["revision"]["head"], "b" * 40)

    def test_process_time_and_output_bounded(self):
        with self.assertRaises(ImportLimit):
            command([sys.executable, "-c", "import time; time.sleep(5)"], timeout=.05, max_bytes=100)
        with self.assertRaises(ImportLimit):
            command([sys.executable, "-c", "print('x'*10000)"], timeout=5, max_bytes=100)

    def test_truncation_is_reported_without_advancing_cursor(self):
        (self.source / "long.md").write_text("# HCCL\n\n" + "Ascend 910B fixture\n" * 200, encoding="utf-8")
        result = self.sync({"path": str(self.source), "limits": {"output_chars": 1000}})
        self.assertEqual(result["status"], "partial")
        self.assertIsNone(result["revision"])
        metadata = self.assert_spans(self.note())
        self.assertFalse(metadata["evidence"]["complete"])
        self.assertTrue(any("budget" in warning for warning in metadata["evidence"]["warnings"]))

    @unittest.skipUnless(HAS_DOCUMENTS, "requires documents test dependencies")
    def test_pdf_page_and_office_archive_bounds_are_partial(self):
        from reportlab.pdfgen import canvas
        from docx import Document
        pdf = canvas.Canvas(str(self.source / "pages.pdf"))
        pdf.drawString(72, 700, "first Ascend page")
        pdf.showPage()
        pdf.drawString(72, 700, "second Ascend page")
        pdf.save()
        result = self.sync({"path": str(self.source / "pages.pdf"), "limits": {"pages": 1}})
        self.assertEqual(result["status"], "partial")
        self.assertIn("first Ascend page", self.note().read_text(encoding="utf-8"))
        self.assertNotIn("second Ascend page", self.note().read_text(encoding="utf-8"))
        document = Document()
        document.add_paragraph("HCCL " * 10000)
        document.save(self.source / "expanded.docx")
        result = self.sync({"path": str(self.source / "expanded.docx"), "limits": {"archive_bytes": 1024}})
        self.assertEqual(result["status"], "partial")
        self.assertIsNone(result["revision"])

    def test_package_has_no_service_or_heavy_import_at_start(self):
        script = "import knowledge_intake,sys; assert not any(x in sys.modules for x in ['mindie_knowledge','docx','pptx','openpyxl','PIL','pypdfium2'])"
        subprocess.check_call([sys.executable, "-c", script])

    @unittest.skipUnless(HAS_DOCUMENTS, "requires documents test dependencies")
    def test_actual_pdf_office_parsers_and_evidence(self):
        from docx import Document
        from pptx import Presentation
        from openpyxl import Workbook
        from reportlab.pdfgen import canvas
        document = Document()
        document.add_heading("HCCL communication", level=1)
        document.add_paragraph("Ascend 910B measured fixture")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "latency", "12 ms"
        document.save(self.source / "hccl.docx")
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = "ACL graph fixture"
        slide.placeholders[1].text = "vLLM Ascend replay"
        presentation.save(self.source / "acl.pptx")
        workbook = Workbook()
        workbook.active.title = "NPU"
        workbook.active.append(["device", "throughput"])
        workbook.active.append(["910B", 123])
        workbook.active["C2"] = "=B2*2"
        workbook.save(self.source / "measurements.xlsx")
        pdf = canvas.Canvas(str(self.source / "report.pdf"))
        pdf.drawString(72, 700, "Ascend graph performance fixture")
        pdf.showPage()
        pdf.drawString(72, 700, "HCCL page two")
        pdf.save()
        first = self.sync()
        self.assertEqual((first["status"], first["converted"]), ("ok", 4), first)
        joined = "\n".join(path.read_text(encoding="utf-8") for path in self.output.glob("*/*.md"))
        for expected in ("12 ms", "vLLM Ascend replay", "B2: 123", "C2: =B2*2", "HCCL page two"):
            self.assertIn(expected, joined)
        for note in self.output.glob("*/*.md"):
            self.assert_spans(note)
        self.assertEqual(self.sync()["converted"], 0)

    def _actual_image_and_pdf_ocr(self, provider, language):
        from knowledge_intake.ocr import capability
        available = capability(provider)
        if language not in available["languages"]:
            self.skipTest(f"optional {provider} OCR/{language}: {available['reason']}; no model download")
        from PIL import Image, ImageDraw, ImageFont
        from reportlab.pdfgen import canvas
        image = Image.new("RGB", (1200, 200), "white")
        draw = ImageDraw.Draw(image)
        try:
            font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 48)
        except OSError:
            font = ImageFont.load_default(size=48)
        draw.text((30, 60), "HCCL graph replay 910B", fill="black", font=font)
        image_path = self.source / "scan.png"
        image.save(image_path)
        pdf = canvas.Canvas(str(self.source / "scan.pdf"), pagesize=(600, 100))
        pdf.drawImage(str(image_path), 0, 0, width=600, height=100)
        pdf.save()
        result = self.sync({"path": str(self.source), "ocr": {"provider": provider, "language": language}})
        self.assertEqual(result["status"], "ok", result)
        self.assertEqual(result["converted"], 2)
        for note in self.output.glob("*/*.md"):
            self.assertIn("HCCL graph replay 910B", note.read_text(encoding="utf-8"))
            metadata = self.assert_spans(note)
            self.assertTrue(any(span["kind"] == "ocr" for span in metadata["evidence"]["spans"]))

    @unittest.skipUnless(HAS_DOCUMENTS, "requires formats test dependencies")
    def test_actual_windows_image_and_scanned_pdf_ocr(self):
        self._actual_image_and_pdf_ocr("windows", "en-US")

    @unittest.skipUnless(HAS_DOCUMENTS, "requires formats test dependencies")
    def test_actual_tesseract_image_and_scanned_pdf_ocr(self):
        self._actual_image_and_pdf_ocr("tesseract", "eng")


if __name__ == "__main__":
    unittest.main()
